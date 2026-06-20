#!/usr/bin/env python3
"""
Sprint 1-B: Importação de clientes e fornecedores
Lê LKL_Base_Completa_v2.xlsx e gera SQL para inserção no banco.
"""
import openpyxl
import re
import sys

XLSX_PATH = '/Users/klebercamara/Desktop/Projeto Grafica LKL/NOVO SISTEMA/LKL_Base_Completa_v2.xlsx'

PJ_KEYWORDS = re.compile(
    r'\b(LTDA|S\.?A\.?|ME|EIRELI|EPP|SS|SLU|COOPERATIVA|COOPE?|ASSOCIACAO|'
    r'FUND[AÇ]|PREFEITURA|SECRETARIA|ESCOLA|COLEGIO|COLEGIO|HOSPITAL|'
    r'CLINICA|SINDICATO|IGREJA|INDUSTRIA|IND\b|COM\b|COMERCIO|SERVICOS?|'
    r'DISTRIBUIDORA|GRAFICA|CONSTRU|LTDA\.|S\/A|CIA)\b',
    re.IGNORECASE
)

def escape_sql(val):
    if val is None:
        return 'NULL'
    s = str(val).strip()
    if not s:
        return 'NULL'
    s = s.replace("'", "''")
    return f"'{s}'"

def clean_digits(val):
    if val is None:
        return None
    return re.sub(r'\D', '', str(val)) or None

def clean_cep(val):
    d = clean_digits(val)
    if d and len(d) == 8:
        return d
    return None

def detect_tipo_pessoa(nome):
    if nome and PJ_KEYWORDS.search(nome):
        return 'PJ'
    return 'PF'

def calc_score(nome, cpf_cnpj, celular, email, cep, logradouro, bairro):
    score = 0
    if nome: score += 20
    if cpf_cnpj: score += 20
    if celular: score += 20
    if email: score += 15
    if cep: score += 10
    if logradouro: score += 10
    if bairro: score += 5
    return score

def clean_phone(val):
    if val is None:
        return None
    # Take only first phone if multiple separated by / or ,
    s = re.split(r'[/,]', str(val))[0]
    d = re.sub(r'[\s\-\(\)\.]', '', s)
    # Truncate to 20 chars max
    return d[:20] if d else None

wb = openpyxl.load_workbook(XLSX_PATH, read_only=True)

lines = []
lines.append('BEGIN;')
lines.append('')

# ── CLIENTES ──────────────────────────────────────────────────────────────────
lines.append('-- CLIENTES')
ws_cli = wb['1. Base Completa']
clientes_count = 0
erros_cli = 0

for row in ws_cli.iter_rows(min_row=5, values_only=True):
    idx, cod, nome, fantasia, email, fone, celular, cep, logradouro, numero, bairro, cidade, uf = row[:13]
    if not nome or not str(nome).strip():
        continue

    nome = str(nome).strip()
    fantasia = str(fantasia).strip() if fantasia else None
    email_v = str(email).strip().lower() if email else None
    fone_v = clean_phone(fone)
    cel_v = clean_phone(celular)
    cep_v = clean_cep(cep)
    end_v = str(logradouro).strip() if logradouro else None
    num_v = str(numero).strip() if numero else None
    bairro_v = str(bairro).strip() if bairro else None
    cidade_v = str(cidade).strip() if cidade else None
    uf_v = str(uf).strip().upper() if uf else None
    tipo = detect_tipo_pessoa(nome)
    score = calc_score(nome, None, cel_v, email_v, cep_v, end_v, bairro_v)
    cod_v = int(cod) if cod else None

    lines.append(
        f"INSERT INTO clientes_lkl (tipo_pessoa, nome, fantasia, email, telefone, celular, cep, "
        f"logradouro, numero, bairro, cidade, uf, canal_origem, score_completude, codigo_sisgraph) VALUES ("
        f"'{tipo}', {escape_sql(nome)}, {escape_sql(fantasia)}, {escape_sql(email_v)}, "
        f"{escape_sql(fone_v)}, {escape_sql(cel_v)}, {escape_sql(cep_v)}, "
        f"{escape_sql(end_v)}, {escape_sql(num_v)}, {escape_sql(bairro_v)}, "
        f"{escape_sql(cidade_v)}, {escape_sql(uf_v)}, 'sisgraph', {score}, "
        f"{'NULL' if cod_v is None else cod_v}"
        f") ON CONFLICT DO NOTHING;"
    )
    clientes_count += 1

lines.append('')

# ── FORNECEDORES ──────────────────────────────────────────────────────────────
lines.append('-- FORNECEDORES')
ws_forn = wb['2. Fornecedores']
forn_count = 0

for row in ws_forn.iter_rows(min_row=5, values_only=True):
    idx, cod, nome, cnpj, contato, ddd, telefone, email, logradouro, cidade, status_raw = row[:11]
    if not nome or not str(nome).strip():
        continue

    nome_v = str(nome).strip()
    cnpj_raw = clean_digits(cnpj)
    # Format CNPJ: 14 digits → XX.XXX.XXX/XXXX-XX
    cnpj_fmt = None
    if cnpj_raw and len(cnpj_raw) == 14:
        cnpj_fmt = f"{cnpj_raw[:2]}.{cnpj_raw[2:5]}.{cnpj_raw[5:8]}/{cnpj_raw[8:12]}-{cnpj_raw[12:]}"
    elif cnpj:
        cnpj_fmt = str(cnpj).strip()

    contato_v = str(contato).strip() if contato else None
    ddd_v = str(ddd).strip() if ddd else None
    tel_v = clean_phone(telefone)
    email_v = str(email).strip().lower() if email else None
    end_v = str(logradouro).strip() if logradouro else None
    cidade_v = str(cidade).strip() if cidade else None
    # Status: ✅ OK = ativo, anything else = inativo
    status_v = 'ativo' if status_raw and '✅' in str(status_raw) else 'inativo'
    cod_v = str(cod).strip() if cod else None

    lines.append(
        f"INSERT INTO fornecedores (nome, cnpj, contato, ddd, telefone, email, logradouro, cidade, "
        f"status, codigo_sisgraph) VALUES ("
        f"{escape_sql(nome_v)}, {escape_sql(cnpj_fmt)}, {escape_sql(contato_v)}, "
        f"{escape_sql(ddd_v)}, {escape_sql(tel_v)}, {escape_sql(email_v)}, "
        f"{escape_sql(end_v)}, {escape_sql(cidade_v)}, '{status_v}', {escape_sql(cod_v)}"
        f") ON CONFLICT DO NOTHING;"
    )
    forn_count += 1

lines.append('')
lines.append('COMMIT;')

output_path = '/Users/klebercamara/LKL/sql/scripts/sprint1b_import.sql'
with open(output_path, 'w', encoding='utf-8') as f:
    f.write('\n'.join(lines))

print(f'✅ SQL gerado: {output_path}')
print(f'   Clientes: {clientes_count}')
print(f'   Fornecedores: {forn_count}')
